#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
compute_sacrebleu.py — official BLEU + chrF (with reproducible signatures) for
every system, using sacrebleu. Run this so the paper can cite sacrebleu.

    pip install sacrebleu openpyxl
    python3 compute_sacrebleu.py

Reads your existing files (edit FILES below if names differ) and prints, per system:
  BLEU, BLEU signature, chrF2, chrF signature.
Also writes sacrebleu_scores.csv.

Scoring pairs:
  • Direct LLMs + the 2 baselines  -> scored vs the DIRECT gold
  • Hybrid (Llama/Mistral/Gemma)    -> scored vs the HYBRID gold
"""
import csv, os
from openpyxl import load_workbook
import sacrebleu
from sacrebleu.metrics import BLEU, CHRF

DIRECT_XLSX = "03_direct_Hindi_to_English_translation_faithfulness.xlsx"
HYBRID_XLSX = "05_hybrid_chunks_to_sentence_3model_BLEU_chrF.xlsx"
GOOGLE_CSV  = "baseline_google_translations.csv"
INDIC_CSV   = "baseline_indictrans2_translations.csv"
def load_direct():
    ws=load_workbook(DIRECT_XLSX,data_only=True)['Analysis']
    gold={}; S={'Llama (direct)':{}, 'Mistral (direct)':{}, 'Gemma (direct)':{}}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        sid,hi,g,lla,lf,li,mis,mf,mi,gem,gf=r[:11]
        if not sid: continue
        gold[sid]=g or ''
        S['Llama (direct)'][sid]=lla or ''; S['Mistral (direct)'][sid]=mis or ''; S['Gemma (direct)'][sid]=gem or ''
    return gold,S
def load_csv(p,c): return {r['sent_id'].strip():(r.get(c) or '') for r in csv.DictReader(open(p,encoding='utf-8-sig'))}
def load_hybrid():
    ws=load_workbook(HYBRID_XLSX,data_only=True)['comparison']
    rows=[r for r in list(ws.iter_rows(values_only=True))[1:] if r[0] and r[3]]
    gold={r[0]:(r[3] or '') for r in rows}
    return gold, {'Llama (hybrid)':{r[0]:(r[4] or '') for r in rows},
                  'Mistral (hybrid)':{r[0]:(r[6] or '') for r in rows},
                  'Gemma (hybrid)':{r[0]:(r[8] or '') for r in rows}}

bleu=BLEU(); chrf=CHRF(word_order=0)   # chrF2 (standard). word_order=2 => chrF++

def score(name, gold, out, sids, rows):
    hyps=[out[s] for s in sids]; refs=[[gold[s] for s in sids]]
    b=bleu.corpus_score(hyps, refs); c=chrf.corpus_score(hyps, refs)
    rows.append([name, round(b.score,1), bleu.get_signature().format(),
                 round(c.score,1), chrf.get_signature().format(), len(sids)])
    print(f"{name:18} BLEU={b.score:5.1f}  chrF2={c.score:5.1f}")

def main():
    rows=[]
    gold,S=load_direct(); sids=list(gold)
    for n,o in S.items(): score(n,gold,o,sids,rows)
    if os.path.exists(GOOGLE_CSV):
        g=load_csv(GOOGLE_CSV,'google_english'); score('Google Translate',gold,{s:g.get(s,'') for s in sids},sids,rows)
    if os.path.exists(INDIC_CSV):
        ind=load_csv(INDIC_CSV,'indictrans2_english'); score('IndicTrans2',gold,{s:ind.get(s,'') for s in sids},sids,rows)
    hg,H=load_hybrid(); hs=list(hg)
    for n,o in H.items(): score(n,hg,o,hs,rows)

    print("\nBLEU signature :", bleu.get_signature().format())
    print("chrF signature :", chrf.get_signature().format())
    with open("sacrebleu_scores.csv","w",newline="",encoding="utf-8") as f:
        w=csv.writer(f); w.writerow(["system","BLEU","BLEU_signature","chrF2","chrF_signature","n"]); w.writerows(rows)
    print("\nSaved sacrebleu_scores.csv")

if __name__=="__main__":
    main()
